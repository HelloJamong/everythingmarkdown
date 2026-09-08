"""Recreate synthetic P0 samples; no downloads or personal documents.

Office writers below are already installed by MarkItDown's selected extras.
This script is a development tool, not an application runtime dependency.
"""

from pathlib import Path
from zipfile import ZIP_DEFLATED, ZipFile

from openpyxl import Workbook
from pptx import Presentation

DEST = Path(__file__).resolve().parents[1] / "tests" / "fixtures"
MARKER = "EverythingMarkdown P0 sample"


def write_pdf(path: Path) -> None:
    stream = f"BT /F1 18 Tf 50 750 Td ({MARKER}) Tj ET\n".encode("ascii")
    objects = [
        b"<< /Type /Catalog /Pages 2 0 R >>",
        b"<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
        b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] "
        b"/Resources << /Font << /F1 4 0 R >> >> /Contents 5 0 R >>",
        b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>",
        f"<< /Length {len(stream)} >>\nstream\n".encode() + stream + b"endstream",
    ]
    data = bytearray(b"%PDF-1.4\n")
    offsets = [0]
    for number, obj in enumerate(objects, 1):
        offsets.append(len(data))
        data.extend(f"{number} 0 obj\n".encode() + obj + b"\nendobj\n")
    start = len(data)
    data.extend(f"xref\n0 {len(offsets)}\n0000000000 65535 f \n".encode())
    for offset in offsets[1:]:
        data.extend(f"{offset:010d} 00000 n \n".encode())
    data.extend(
        f"trailer\n<< /Size {len(offsets)} /Root 1 0 R >>\n"
        f"startxref\n{start}\n%%EOF\n".encode()
    )
    path.write_bytes(data)


def main() -> None:
    DEST.mkdir(parents=True, exist_ok=True)
    (DEST / "sample.txt").write_text(f"{MARKER}\n한글 확인\n", encoding="utf-8")
    (DEST / "sample.csv").write_text(f"title,value\n{MARKER},42\n", encoding="utf-8")
    (DEST / "sample.html").write_text(
        f'<!doctype html><html><head><meta charset="utf-8"></head>'
        f"<body><h1>{MARKER}</h1><p>한글 확인</p></body></html>\n",
        encoding="utf-8",
    )
    write_pdf(DEST / "sample.pdf")
    with ZipFile(DEST / "sample.docx", "w", ZIP_DEFLATED) as archive:
        archive.writestr(
            "[Content_Types].xml",
            '<Types xmlns="http://schemas.openxmlformats.org/package/2006/content-types">'
            '<Default Extension="rels" ContentType="application/vnd.openxmlformats-package.relationships+xml"/>'
            '<Override PartName="/word/document.xml" ContentType="application/vnd.openxmlformats-officedocument.wordprocessingml.document.main+xml"/>'
            '</Types>',
        )
        archive.writestr(
            "_rels/.rels",
            '<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">'
            '<Relationship Id="rId1" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/officeDocument" Target="word/document.xml"/>'
            '</Relationships>',
        )
        archive.writestr(
            "word/document.xml",
            '<w:document xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main">'
            f'<w:body><w:p><w:r><w:t>{MARKER}</w:t></w:r></w:p></w:body></w:document>',
        )
    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[0])
    slide.shapes.title.text = MARKER
    presentation.save(DEST / "sample.pptx")
    workbook = Workbook()
    workbook.active.append(["title", "value"])
    workbook.active.append([MARKER, 42])
    workbook.save(DEST / "sample.xlsx")


if __name__ == "__main__":
    main()
